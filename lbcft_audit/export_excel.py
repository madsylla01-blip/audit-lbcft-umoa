"""Export des résultats dans un classeur Excel de feuilles de travail (un onglet par test)."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .moteur import ResultatsAudit
from .pseudonymisation import pseudonymiser_table

ENTETE = PatternFill("solid", fgColor="DCE6F2")
TITRE = Font(bold=True, size=14, color="0B0B0B")
GRAS = Font(bold=True)
GRIS = Font(color="52514E", italic=True)
BORDURE = Border(bottom=Side(style="thin", color="C3C2B7"))
FORMAT_MONTANT = "#,##0"
FORMAT_DATE = "dd/mm/yyyy"
FORMAT_DATE_HEURE = "dd/mm/yyyy hh:mm"


def _valeur(v):
    if v is None or (not isinstance(v, (list, dict, set)) and pd.isna(v)):
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    if hasattr(v, "item"):
        return v.item()
    return v


def _tableau(ws, table: pd.DataFrame, ligne: int) -> int:
    """Écrit un DataFrame à partir de `ligne` ; renvoie la ligne suivante."""
    for j, colonne in enumerate(table.columns, start=1):
        cellule = ws.cell(row=ligne, column=j, value=str(colonne))
        cellule.font, cellule.fill, cellule.border = GRAS, ENTETE, BORDURE
        cellule.alignment = Alignment(wrap_text=True, vertical="top")
    for i, enregistrement in enumerate(table.itertuples(index=False), start=ligne + 1):
        for j, v in enumerate(enregistrement, start=1):
            cellule = ws.cell(row=i, column=j, value=_valeur(v))
            nom = str(table.columns[j - 1])
            if isinstance(cellule.value, datetime):
                cellule.number_format = (
                    FORMAT_DATE_HEURE
                    if nom == "date_operation" and (cellule.value.hour or cellule.value.minute)
                    else FORMAT_DATE
                )
            elif isinstance(cellule.value, float) and (nom.startswith("taux") or nom in ("precision", "rappel")):
                cellule.number_format = "0.0%"
            elif isinstance(cellule.value, (int, float)) and any(
                k in nom for k in ("montant", "credit", "debit", "revenu", "cumul", "plafond")
            ):
                cellule.number_format = FORMAT_MONTANT
    if len(table.columns):
        fin = ligne + max(len(table), 1)
        ws.auto_filter.ref = f"A{ligne}:{get_column_letter(len(table.columns))}{fin}"
    for j, colonne in enumerate(table.columns, start=1):
        valeurs = [str(colonne)] + [str(x) for x in table.iloc[:200, j - 1].tolist()]
        ws.column_dimensions[get_column_letter(j)].width = min(max(len(x) for x in valeurs) + 2, 60)
    return ligne + len(table) + 2


def _param(valeur) -> str:
    if isinstance(valeur, (int, float)) and not isinstance(valeur, bool) and abs(valeur) >= 1000:
        return f"{valeur:,.0f}".replace(",", " ")
    return str(valeur)


def _entete_feuille(ws, titre: str, lignes: list[tuple[str, str]]) -> int:
    ws["A1"] = titre
    ws["A1"].font = TITRE
    for i, (cle, valeur) in enumerate(lignes, start=2):
        ws.cell(row=i, column=1, value=cle).font = GRAS
        ws.cell(row=i, column=2, value=valeur).alignment = Alignment(wrap_text=False)
    return len(lignes) + 3


def exporter_excel(resultats: ResultatsAudit, chemin: str | Path) -> Path:
    p, d = resultats.parametres, resultats.donnees
    pseudo = bool(p["pseudonymisation"]["active"])
    chemin = Path(chemin)
    chemin.parent.mkdir(parents=True, exist_ok=True)
    classeur = Workbook()

    # --- Synthèse
    ws = classeur.active
    ws.title = "Synthèse"
    ligne = _entete_feuille(
        ws,
        "Audit du dispositif LBC/FT — synthèse des tests",
        [
            ("Entité auditée", p.banque + " — données 100 % fictives"),
            ("Période testée", f"{p.periode_debut:%d/%m/%Y} au {p.periode_fin:%d/%m/%Y}"),
            ("Date d'arrêté des données", f"{p.date_arrete:%d/%m/%Y}"),
            ("Exécution", f"{datetime.now():%d/%m/%Y %H:%M} — script lbcft_audit (reproductible)"),
            ("Statut", "Projet de feuilles de travail — exceptions à qualifier sur pièces avant toute conclusion"),
        ],
    )
    synthese = resultats.synthese().drop(columns=["unite"])
    _tableau(ws, synthese, ligne)
    ws.freeze_panes = ws.cell(row=ligne + 1, column=1)

    # --- Méthodologie
    ws = classeur.create_sheet("Méthodologie")
    ligne = _entete_feuille(
        ws,
        "Méthodologie et paramètres",
        [
            ("Démarche", "Contrôles d'intégrité, puis réexécution exhaustive des tests sur 100 % de la population."),
            ("Qualification", "Chaque exception est une anomalie à expliquer : revue sur pièces avant conclusion."),
            (
                "Soupçon",
                "L'audit ne déclare pas : les cas sont transmis à la conformité, seule habilitée à déclarer à la CENTIF.",
            ),
            ("Pseudonymisation", "Active" if pseudo else "Inactive (données fictives)"),
            ("Paramètres", "config/parametres.yaml — seuils réglementaires sourcés, hypothèses d'audit signalées"),
        ],
    )
    params = [
        {"test": t.code, "titre": t.titre, "parametre": k, "valeur": _param(v), "references": t.references}
        for t in resultats.tests.values()
        for k, v in t.parametres.items()
    ]
    ligne = _tableau(ws, pd.DataFrame(params), ligne)
    sources = pd.DataFrame([{"table": nom, "fichier": f"{nom}.csv", "lignes": len(t)} for nom, t in d.tables().items()])
    _tableau(ws, sources, ligne)

    # --- Intégrité
    ws = classeur.create_sheet("Intégrité")
    ligne = _entete_feuille(
        ws,
        "Contrôles d'intégrité des données",
        [("Rappel", "Sans rapprochement des totaux avec la comptabilité, les résultats restent provisoires.")],
    )
    _tableau(ws, resultats.integrite, ligne)

    # --- Une feuille de travail par test
    for t in resultats.tests.values():
        ws = classeur.create_sheet(t.code)
        ligne = _entete_feuille(
            ws,
            f"{t.code} — {t.titre}",
            [
                ("Objectif", t.objectif),
                ("Références", t.references),
                ("Paramètres", " ; ".join(f"{k} = {_param(v)}" for k, v in t.parametres.items())),
                ("Population testée", f"{t.nb_population} {t.unite if t.unite != 'correspondances' else 'clients'}"),
                (
                    "Exceptions",
                    f"{t.nb_exceptions} {t.unite}"
                    + (f" — {t.montant_exceptions:,.0f} FCFA".replace(",", " ") if t.montant_exceptions else ""),
                ),
                ("Commentaire", t.commentaire or "—"),
                ("Conclusion", "À compléter par l'auditeur après revue sur pièces"),
                ("Préparé par / date", ""),
                ("Revu par / date", ""),
            ],
        )
        exceptions = pseudonymiser_table(t.exceptions) if pseudo else t.exceptions
        if exceptions.empty:
            ws.cell(row=ligne, column=1, value="Aucune exception.").font = GRIS
        else:
            _tableau(ws, exceptions, ligne)
            ws.freeze_panes = ws.cell(row=ligne + 1, column=1)

    # --- Annexes utiles à l'analyse
    annexes = [
        ("OPE-05", "ecarts_liste", "Pays - écarts liste"),
        ("OPE-05", "par_pays", "Pays - flux"),
        ("COU-01", "couverture_par_test", "Couverture par test"),
        ("COU-01", "traitement_banque", "Couverture - traitement"),
        ("DEC-01", "par_agence", "DTE par agence"),
        ("KYC-01", "par_agence", "KYC par agence"),
        ("ALE-01", "par_analyste", "Alertes par analyste"),
    ]
    for code, cle, nom in annexes:
        ws = classeur.create_sheet(nom[:31])
        ligne = _entete_feuille(ws, f"{code} — annexe : {nom}", [])
        _tableau(ws, resultats[code].annexes[cle], ligne)
    ws = classeur.create_sheet("Priorisation")
    ligne = _entete_feuille(
        ws,
        "Priorisation des revues sur pièces (score de l'auditeur)",
        [("Lecture", "Le score oriente l'ordre des revues ; il ne qualifie pas un soupçon.")],
    )
    _tableau(ws, resultats.priorisation, ligne)

    classeur.save(chemin)
    return chemin
